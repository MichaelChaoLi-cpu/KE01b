#!/usr/bin/env python3
"""Network and Simulation Descriptive Summary.

Plan: Summarize junction-defined roads, demand units, emergency facilities,
network connectors, failure calibration, and executed simulation scale.
Framework: AnaSOP Section 6.1, Section 6.6, and Analytical Workflow steps 1-3.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
PROCESSED = ROOT / "data" / "processed"
FORMAL = ROOT / "data" / "exp" / "monte_carlo_length_weighted_full_1000"
ROAD_CONSEQUENCE = ROOT / "data" / "exp" / "road_section_accessibility_consequence"
SPEED = ROOT / "data" / "exp" / "monte_carlo_speed_sensitivity_1000"
OUTPUT = (
    ROOT
    / "data"
    / "results"
    / "tables"
    / "Table_network_and_simulation_descriptive_summary.xlsx"
)
STAT_COLUMNS = [
    "Count",
    "Total",
    "Mean",
    "Standard Deviation",
    "P05",
    "P50",
    "P95",
]


def numeric(values: Iterable[object]) -> np.ndarray:
    series = pd.to_numeric(pd.Series(values), errors="coerce").dropna()
    return series.to_numpy(dtype=float)


def distribution_row(
    category: str,
    indicator: str,
    unit: str,
    values: Iterable[object],
    *,
    total: float | None = None,
) -> dict[str, object]:
    array = numeric(values)
    if not len(array):
        raise ValueError(f"No numeric values for {indicator}")
    return {
        "Category": category,
        "Indicator": indicator,
        "Unit": unit,
        "Count": int(len(array)),
        "Total": None if total is None else float(total),
        "Mean": float(array.mean()),
        "Standard Deviation": float(array.std(ddof=1)) if len(array) > 1 else 0.0,
        "P05": float(np.quantile(array, 0.05)),
        "P50": float(np.quantile(array, 0.50)),
        "P95": float(np.quantile(array, 0.95)),
    }


def count_row(
    category: str,
    indicator: str,
    unit: str,
    count: int,
    *,
    total: float | None = None,
    mean: float | None = None,
) -> dict[str, object]:
    return {
        "Category": category,
        "Indicator": indicator,
        "Unit": unit,
        "Count": int(count),
        "Total": total,
        "Mean": mean,
        "Standard Deviation": None,
        "P05": None,
        "P50": None,
        "P95": None,
    }


def accepted_connector_rows(
    category: str,
    label: str,
    frame: pd.DataFrame,
    eligibility: pd.Series | None = None,
) -> list[dict[str, object]]:
    eligible = pd.Series(True, index=frame.index) if eligibility is None else eligibility
    eligible = eligible.fillna(False)
    accepted = eligible & frame["Network Snap Accepted"].fillna(False)
    eligible_count = int(eligible.sum())
    accepted_count = int(accepted.sum())
    rows = [
        count_row(
            category,
            f"Accepted {label} connectors",
            "connectors; acceptance share in Mean",
            accepted_count,
            total=eligible_count,
            mean=accepted_count / eligible_count if eligible_count else np.nan,
        )
    ]
    if accepted_count:
        rows.append(
            distribution_row(
                category,
                f"Accepted {label} connector snap distance",
                "metres",
                frame.loc[accepted, "Network Snap Distance (m)"],
            )
        )
    return rows


def style_workbook(path: Path, data_rows: int) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook["Descriptive Summary"]
    navy = "1F4E78"
    blue = "D9EAF7"
    light_blue = "EAF3F8"
    white = "FFFFFF"
    grey = "D9E1E5"
    thin_grey = Side(style="thin", color="B7C3CA")

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    title = sheet.cell(1, 1)
    title.value = "Network and Simulation Descriptive Summary"
    title.font = Font(name="Arial", size=14, bold=True, color=white)
    title.fill = PatternFill("solid", fgColor=navy)
    title.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 24

    header_row = 3
    for cell in sheet[header_row]:
        cell.font = Font(name="Arial", size=9, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin_grey, bottom=thin_grey, left=thin_grey, right=thin_grey)
    sheet.row_dimensions[header_row].height = 31

    category_fills = {
        "Road network": "DCE6F1",
        "Population and connectors": "E2F0D9",
        "Emergency facilities": "FFF2CC",
        "Simulation design and execution": "FCE4D6",
    }
    first_data_row = 4
    last_data_row = first_data_row + data_rows - 1
    previous_category = None
    for row in range(first_data_row, last_data_row + 1):
        category_value = sheet.cell(row, 1).value
        category = category_value if category_value is not None else previous_category
        fill = PatternFill("solid", fgColor=category_fills.get(category, blue))
        for column in range(1, 11):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Arial", size=8.5)
            cell.border = Border(
                top=thin_grey,
                bottom=thin_grey,
                left=thin_grey,
                right=thin_grey,
            )
            cell.alignment = Alignment(
                horizontal="left" if column <= 3 else "right",
                vertical="center",
                wrap_text=column <= 3,
            )
        sheet.cell(row, 1).fill = fill
        sheet.cell(row, 1).font = Font(name="Arial", size=8.5, bold=True)
        if category_value is not None and row > first_data_row:
            for column in range(1, 11):
                sheet.cell(row, column).border = Border(
                    top=Side(style="medium", color="7F8C8D"),
                    bottom=thin_grey,
                    left=thin_grey,
                    right=thin_grey,
                )
        previous_category = category

    for row in range(first_data_row, last_data_row + 1):
        sheet.cell(row, 4).number_format = "#,##0"
        for column in range(5, 11):
            sheet.cell(row, column).number_format = "#,##0.000"

        unit = str(sheet.cell(row, 3).value or "")
        indicator = str(sheet.cell(row, 2).value or "")
        integer_total_units = {
            "sections",
            "edges",
            "nodes",
            "components",
            "people",
            "connectors; acceptance share in Mean",
            "disclosure groups",
            "fire-service facilities",
            "hospitals",
            "beds",
            "replicates",
            "network states",
            "severity outcomes",
        }
        if unit in integer_total_units:
            sheet.cell(row, 5).number_format = "#,##0"
        if indicator == "Simulation replicates per severity":
            for column in range(6, 11):
                sheet.cell(row, column).number_format = "#,##0"

    note_row = last_data_row + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
    note = sheet.cell(note_row, 1)
    note.value = (
        "Note: Statistics are reported where applicable; blank cells indicate that a "
        "statistic is not meaningful for the indicator. For connector rows, Count is "
        "accepted connectors, Total is eligible connectors, and Mean is the acceptance "
        "share. For road-length rows, Total is kilometres, while distribution statistics "
        "are metres."
    )
    note.font = Font(name="Arial", size=9, italic=True, color="37474F")
    note.fill = PatternFill("solid", fgColor=light_blue)
    note.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[note_row].height = 42

    widths = [
        27,
        45,
        31,
        14,
        17,
        17,
        19,
        15,
        15,
        15,
    ]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "D4"
    sheet.auto_filter.ref = f"A3:J{last_data_row}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(path)


def main() -> None:
    sections = pd.read_parquet(PROCESSED / "kumamoto_road_sections_preprocessed.parquet")
    edges = pd.read_parquet(PROCESSED / "kumamoto_routable_road_edges_preprocessed.parquet")
    nodes = pd.read_parquet(PROCESSED / "kumamoto_routable_road_nodes_preprocessed.parquet")
    mesh = pd.read_parquet(
        PROCESSED / "kumamoto_population_mesh_network_access_preprocessed.parquet"
    )
    older = pd.read_parquet(
        PROCESSED / "kumamoto_population_group_network_access_preprocessed.parquet"
    )
    dispatch = pd.read_parquet(
        PROCESSED / "kumamoto_dispatch_base_network_access_preprocessed.parquet"
    )
    hospitals = pd.read_parquet(
        PROCESSED / "kumamoto_hospital_network_access_preprocessed.parquet"
    )
    replicate_metrics = pd.read_parquet(FORMAL / "replicate_metrics.parquet")
    convergence = pd.read_parquet(FORMAL / "convergence_summary.parquet")
    road_consequence = pd.read_parquet(
        ROAD_CONSEQUENCE / "road_section_accessibility_consequence.parquet"
    )
    formal_summary = json.loads((FORMAL / "experiment_summary.json").read_text())
    speed_summary = json.loads((SPEED / "experiment_summary.json").read_text())

    eligible_sections = sections["Road Available"].fillna(False) & sections[
        "Network Analysis Eligible"
    ].fillna(False)
    eligible_edges = edges["Road Available"].fillna(False) & edges[
        "Network Analysis Eligible"
    ].fillna(False)
    eligible_nodes = nodes["Network Analysis Eligible"].fillna(False)
    road_section_frame = sections.loc[eligible_sections]
    road_edge_frame = edges.loc[eligible_edges]

    rows: list[dict[str, object]] = []
    rows.extend(
        [
            count_row(
                "Road network",
                "Eligible junction-to-junction road sections",
                "sections",
                len(road_section_frame),
            ),
            distribution_row(
                "Road network",
                "Road section length",
                "Total: km; distribution: metres",
                road_section_frame["Road Section Length (m)"],
                total=float(road_section_frame["Road Section Length (m)"].sum() / 1000),
            ),
            distribution_row(
                "Road network",
                "Internal road edges per road section",
                "edges",
                road_section_frame["Road Edge Count"],
                total=float(road_section_frame["Road Edge Count"].sum()),
            ),
            distribution_row(
                "Road network",
                "Baseline road-section travel time",
                "minutes",
                road_section_frame["Baseline Section Travel Time (min)"],
            ),
            distribution_row(
                "Road network",
                "Assumed road-section speed",
                "km/h",
                road_section_frame["Assumed Speed (km/h)"],
            ),
            count_row(
                "Road network",
                "Eligible routable road edges",
                "edges",
                len(road_edge_frame),
            ),
            distribution_row(
                "Road network",
                "Routable road-edge length",
                "Total: km; distribution: metres",
                road_edge_frame["Road Length (m)"],
                total=float(road_edge_frame["Road Length (m)"].sum() / 1000),
            ),
            count_row(
                "Road network",
                "Eligible physical road nodes",
                "nodes",
                int(eligible_nodes.sum()),
            ),
            count_row(
                "Road network",
                "Eligible network components",
                "components",
                int(road_section_frame["Network Component ID"].nunique()),
            ),
            count_row(
                "Road network",
                "Emergency-route member sections",
                "sections",
                int(
                    road_section_frame["Emergency Route Membership"]
                    .fillna("None")
                    .ne("None")
                    .sum()
                ),
                total=len(road_section_frame),
                mean=float(
                    road_section_frame["Emergency Route Membership"]
                    .fillna("None")
                    .ne("None")
                    .mean()
                ),
            ),
            count_row(
                "Road network",
                "Landslide-hazard-exposed sections",
                "sections",
                int(
                    road_section_frame["Hazard Exposure Class"]
                    .fillna("None")
                    .ne("None")
                    .sum()
                ),
                total=len(road_section_frame),
                mean=float(
                    road_section_frame["Hazard Exposure Class"]
                    .fillna("None")
                    .ne("None")
                    .mean()
                ),
            ),
        ]
    )

    rows.extend(
        [
            count_row(
                "Population and connectors",
                "Population mesh demand units",
                "mesh units",
                len(mesh),
            ),
            distribution_row(
                "Population and connectors",
                "Total population per mesh demand unit",
                "people",
                mesh["Total Population"],
                total=float(mesh["Total Population"].sum()),
            ),
        ]
    )
    rows.extend(
        accepted_connector_rows(
            "Population and connectors", "population mesh", mesh
        )
    )
    rows.extend(
        [
            count_row(
                "Population and connectors",
                "Older-population disclosure groups",
                "disclosure groups",
                len(older),
            ),
            distribution_row(
                "Population and connectors",
                "Population age 65+ per disclosure group",
                "people",
                older["Population Age 65+"],
                total=float(older["Population Age 65+"].sum()),
            ),
        ]
    )
    rows.extend(
        accepted_connector_rows(
            "Population and connectors", "older-population group", older
        )
    )

    candidate_dispatch = dispatch["Candidate Dispatch Base"].fillna(False)
    eligible_hospitals = hospitals["Eligible Emergency Hospital"].fillna(False)
    rows.append(
        count_row(
            "Emergency facilities",
            "Candidate ambulance dispatch bases",
            "fire-service facilities",
            int(candidate_dispatch.sum()),
        )
    )
    rows.extend(
        accepted_connector_rows(
            "Emergency facilities",
            "dispatch-base",
            dispatch,
            candidate_dispatch,
        )
    )
    rows.append(
        count_row(
            "Emergency facilities",
            "Eligible emergency hospitals",
            "hospitals",
            int(eligible_hospitals.sum()),
        )
    )
    rows.extend(
        accepted_connector_rows(
            "Emergency facilities",
            "hospital",
            hospitals,
            eligible_hospitals,
        )
    )
    reported_capacity = eligible_hospitals & hospitals["Hospital Capacity Weight"].notna()
    rows.extend(
        [
            count_row(
                "Emergency facilities",
                "Eligible hospitals with reported capacity",
                "hospitals",
                int(reported_capacity.sum()),
                total=int(eligible_hospitals.sum()),
                mean=float(reported_capacity.sum() / eligible_hospitals.sum()),
            ),
            distribution_row(
                "Emergency facilities",
                "Reported hospital capacity weight",
                "beds",
                hospitals.loc[reported_capacity, "Hospital Capacity Weight"],
                total=float(
                    hospitals.loc[reported_capacity, "Hospital Capacity Weight"].sum()
                ),
            ),
        ]
    )

    targets = np.array(
        formal_summary["design"]["expected_failed_road_length_shares"], dtype=float
    )
    final_convergence = convergence.loc[
        convergence["Simulation Replicate Checkpoint"].eq(1000)
    ]
    rows.extend(
        [
            distribution_row(
                "Simulation design and execution",
                "Expected failed road-length targets",
                "share",
                targets,
            ),
            count_row(
                "Simulation design and execution",
                "Main failure-severity levels",
                "severity levels",
                3,
            ),
            count_row(
                "Simulation design and execution",
                "Stress failure-severity levels",
                "severity levels",
                1,
            ),
            distribution_row(
                "Simulation design and execution",
                "Simulation replicates per severity",
                "replicates",
                [formal_summary["design"]["replicates_per_severity"]] * len(targets),
                total=float(formal_summary["design"]["network_states"]),
            ),
            count_row(
                "Simulation design and execution",
                "Formal road-failure network states",
                "network states",
                formal_summary["design"]["network_states"],
            ),
            distribution_row(
                "Simulation design and execution",
                "Realized failed road-length share",
                "share",
                replicate_metrics["Realized Failed Road Length Share"],
            ),
            distribution_row(
                "Simulation design and execution",
                "Failed road sections per network state",
                "sections",
                replicate_metrics["Failed Road Sections"],
            ),
            distribution_row(
                "Simulation design and execution",
                "Routing time per network state",
                "seconds",
                replicate_metrics["Routing Seconds"],
                total=float(replicate_metrics["Routing Seconds"].sum()),
            ),
            count_row(
                "Simulation design and execution",
                "Severity outcomes stable at 1,000 replicates",
                "severity outcomes",
                int(final_convergence["Monte Carlo Convergence Status"].eq("Stable").sum()),
                total=len(final_convergence),
                mean=float(final_convergence["Monte Carlo Convergence Status"].eq("Stable").mean()),
            ),
            distribution_row(
                "Simulation design and execution",
                "Uniform assumed-speed sensitivity multipliers",
                "multiplier",
                speed_summary["assumed_speed_multipliers"],
            ),
            count_row(
                "Simulation design and execution",
                "Speed-sensitivity road states replayed",
                "network states",
                speed_summary["network_states_replayed"],
            ),
            count_row(
                "Simulation design and execution",
                "Road sections in all-section consequence analysis",
                "sections",
                len(road_consequence),
            ),
            count_row(
                "Simulation design and execution",
                "Road sections individually rerouted",
                "sections",
                int(road_consequence["Candidate for Rerouting"].sum()),
                total=len(road_consequence),
                mean=float(road_consequence["Candidate for Rerouting"].mean()),
            ),
        ]
    )

    table = pd.DataFrame(rows)
    table = table.set_index(["Category", "Indicator", "Unit"])[STAT_COLUMNS]
    table = table.where(pd.notna(table), None)
    if len(table) < 30 or len(table) > 45:
        raise RuntimeError(f"Unexpected descriptive-table row count: {len(table)}")
    if len(road_section_frame) != 343844:
        raise RuntimeError("Eligible road-section count changed")
    if formal_summary["design"]["network_states"] != 4000:
        raise RuntimeError("Formal network-state count changed")
    if not final_convergence["Monte Carlo Convergence Status"].eq("Stable").all():
        raise RuntimeError("A final severity outcome is not stable")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Descriptive Summary", startrow=2, index=True)
    style_workbook(OUTPUT, len(table))
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Rows: {len(table)}; statistic columns: {len(STAT_COLUMNS)}")
    print(table.to_string())


if __name__ == "__main__":
    main()
