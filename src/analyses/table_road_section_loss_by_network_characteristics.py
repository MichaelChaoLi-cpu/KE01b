#!/usr/bin/env python3
"""Road-Section Loss by Network Characteristics.

Plan: Summarize full-network potential 30-minute access loss and scenario-specific
expected risk by road category, emergency-route status, and hazard exposure.
Framework: AnaSOP Sections 5.3, 6.5, and Analytical Workflow step 7.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
CONSEQUENCE = (
    ROOT
    / "data"
    / "exp"
    / "road_section_accessibility_consequence"
    / "road_section_accessibility_consequence.parquet"
)
FAILURE = (
    ROOT
    / "data"
    / "exp"
    / "monte_carlo_length_weighted_full_1000"
    / "nested_failure_realization_replicate_001.parquet"
)
OUTPUT = (
    ROOT
    / "data"
    / "results"
    / "tables"
    / "Table_road_section_loss_by_network_characteristics.xlsx"
)
SEVERITIES = (1, 3, 5, 10)
POTENTIAL_LOSS = "Potential Access Loss 30 Minutes"
TABLE_COLUMNS = [
    "Network Characteristic Group",
    "Expected Failed Road Length (%)",
    "Road Section Count",
    "Road Length (km)",
    "Zero Potential-Loss Share (%)",
    "Mean Potential Access Loss",
    "Potential Access Loss P50",
    "Potential Access Loss P90",
    "Potential Access Loss P95",
    "Potential Access Loss P99",
    "Maximum Potential Access Loss",
    "Mean Expected Risk",
]


def load_analysis() -> pd.DataFrame:
    consequence = pd.read_parquet(CONSEQUENCE)
    failure_columns = ["Road Section ID"] + [
        f"Section Failure Probability {severity}%" for severity in SEVERITIES
    ]
    failure = pd.read_parquet(FAILURE, columns=failure_columns)
    if consequence["Road Section ID"].duplicated().any():
        raise RuntimeError("Road Section ID is not unique in consequence results")
    if failure["Road Section ID"].duplicated().any():
        raise RuntimeError("Road Section ID is not unique in failure results")
    frame = consequence.merge(
        failure, on="Road Section ID", how="left", validate="one_to_one"
    )
    if len(frame) != 343844:
        raise RuntimeError(f"Expected 343,844 road sections; got {len(frame)}")
    required = [
        "Road Section Length (m)",
        "Road Category",
        "Emergency Route Membership",
        "Hazard Exposure Class",
        POTENTIAL_LOSS,
    ] + [f"Section Failure Probability {severity}%" for severity in SEVERITIES]
    if frame[required].isna().any().any():
        raise RuntimeError("Required road-consequence values are missing")
    if (frame[POTENTIAL_LOSS] < 0).any():
        raise RuntimeError("Potential access loss is negative")
    for severity in SEVERITIES:
        probability = frame[f"Section Failure Probability {severity}%"]
        if not probability.between(0, 1).all():
            raise RuntimeError(f"Failure probability is invalid at {severity}%")
        frame[f"Expected Risk {severity}%"] = probability * frame[POTENTIAL_LOSS]
        if severity in (1, 3, 5):
            source = frame[f"Expected Risk 30 Minutes {severity} Percent"]
            if not np.allclose(
                frame[f"Expected Risk {severity}%"], source, atol=1e-9
            ):
                raise RuntimeError(
                    f"Expected-risk reconstruction failed at {severity}%"
                )
    probabilities = frame[
        [f"Section Failure Probability {severity}%" for severity in SEVERITIES]
    ].to_numpy(float)
    if (np.diff(probabilities, axis=1) < -1e-12).any():
        raise RuntimeError("Road-section failure probabilities are not nested")
    return frame


def grouping_definitions(frame: pd.DataFrame) -> list[tuple[str, pd.Series]]:
    groups: list[tuple[str, pd.Series]] = [
        ("All eligible road sections", pd.Series(True, index=frame.index))
    ]
    group_columns = [
        ("Road Category", "Road Category"),
        ("Emergency Route Membership", "Emergency Route Membership"),
        ("Hazard Exposure Class", "Hazard Exposure Class"),
    ]
    expected_levels = {
        "Road Category": 5,
        "Emergency Route Membership": 3,
        "Hazard Exposure Class": 3,
    }
    for label, column in group_columns:
        levels = sorted(frame[column].astype(str).unique())
        if len(levels) != expected_levels[column]:
            raise RuntimeError(
                f"Unexpected {column} level count: {len(levels)}"
            )
        for level in levels:
            groups.append(
                (f"{label} — {level}", frame[column].astype(str).eq(level))
            )
    if len(groups) != 12:
        raise RuntimeError(f"Expected 12 network-characteristic groups; got {len(groups)}")
    return groups


def summarize(frame: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, float | int | str]] = []
    for group_label, mask in grouping_definitions(frame):
        subset = frame.loc[mask]
        potential = subset[POTENTIAL_LOSS].to_numpy(float)
        if len(subset) == 0:
            raise RuntimeError(f"Empty network-characteristic group: {group_label}")
        shared = {
            "Network Characteristic Group": group_label,
            "Road Section Count": int(len(subset)),
            "Road Length (km)": float(
                subset["Road Section Length (m)"].sum() / 1000.0
            ),
            "Zero Potential-Loss Share (%)": float(100.0 * np.mean(potential == 0)),
            "Mean Potential Access Loss": float(potential.mean()),
            "Potential Access Loss P50": float(np.quantile(potential, 0.50)),
            "Potential Access Loss P90": float(np.quantile(potential, 0.90)),
            "Potential Access Loss P95": float(np.quantile(potential, 0.95)),
            "Potential Access Loss P99": float(np.quantile(potential, 0.99)),
            "Maximum Potential Access Loss": float(potential.max()),
        }
        for severity in SEVERITIES:
            rows.append(
                {
                    **shared,
                    "Expected Failed Road Length (%)": severity,
                    "Mean Expected Risk": float(
                        subset[f"Expected Risk {severity}%"].mean()
                    ),
                }
            )
    table = pd.DataFrame(rows, columns=TABLE_COLUMNS)
    if table.shape != (48, 12):
        raise RuntimeError(f"Unexpected road-characteristics table shape: {table.shape}")
    if table.isna().any().any():
        raise RuntimeError("Road-characteristics table contains missing values")
    for _, group in table.groupby("Network Characteristic Group", sort=False):
        if not group["Road Section Count"].nunique() == 1:
            raise RuntimeError("Road-section count varies within a group")
        if (np.diff(group["Mean Expected Risk"].to_numpy(float)) < -1e-12).any():
            raise RuntimeError("Mean expected risk decreases with severity")
    return table


def style_workbook(path: Path, data_rows: int) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook["Network Characteristics"]
    navy = "1F4E78"
    main_fill = "F4F8FB"
    main_alt_fill = "E8F1F7"
    stress_fill = "FCE4D6"
    note_fill = "EAF3F8"
    white = "FFFFFF"
    thin_grey = Side(style="thin", color="B7C3CA")
    medium_grey = Side(style="medium", color="7F8C8D")

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
    title = sheet.cell(1, 1)
    title.value = "Road-Section Loss by Network Characteristics"
    title.font = Font(name="Arial", size=14, bold=True, color=white)
    title.fill = PatternFill("solid", fgColor=navy)
    title.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 24

    header_row = 3
    for cell in sheet[header_row]:
        cell.font = Font(name="Arial", size=8.5, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=thin_grey,
            bottom=thin_grey,
            left=thin_grey,
            right=thin_grey,
        )
    sheet.row_dimensions[header_row].height = 70

    first_data_row = 4
    last_data_row = first_data_row + data_rows - 1
    group_number = -1
    previous_group: str | None = None
    for row in range(first_data_row, last_data_row + 1):
        group_label = str(sheet.cell(row, 1).value)
        severity = float(sheet.cell(row, 2).value)
        starts_group = group_label != previous_group
        if starts_group:
            group_number += 1
        if np.isclose(severity, 10.0):
            fill = PatternFill("solid", fgColor=stress_fill)
        else:
            fill = PatternFill(
                "solid", fgColor=main_fill if group_number % 2 == 0 else main_alt_fill
            )
        top_side = medium_grey if starts_group and row > first_data_row else thin_grey
        for column in range(1, 13):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Arial", size=8.5)
            cell.fill = fill
            cell.border = Border(
                top=top_side,
                bottom=thin_grey,
                left=thin_grey,
                right=thin_grey,
            )
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "right",
                vertical="center",
                wrap_text=column == 1,
            )
        sheet.cell(row, 2).number_format = '0"%"'
        sheet.cell(row, 3).number_format = "#,##0"
        sheet.cell(row, 4).number_format = "#,##0.0"
        sheet.cell(row, 5).number_format = '0.0"%"'
        for column in range(6, 12):
            sheet.cell(row, column).number_format = "#,##0.0"
        sheet.cell(row, 11).number_format = "#,##0"
        sheet.cell(row, 12).number_format = "#,##0.000"
        previous_group = group_label

    note_row = last_data_row + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=12)
    note = sheet.cell(note_row, 1)
    note.value = (
        "Note: Potential access loss is the population losing 30-minute emergency access "
        "when one junction-to-junction road section is removed and all other sections "
        "remain available. It is severity-independent and repeats across scenario rows. "
        "Mean expected risk equals section failure probability times potential access loss. "
        "The 1%, 3%, and 5% states are main scenarios; 10% is an extreme stress scenario. "
        "Characteristic groups overlap and are not additive. None denotes no emergency-route "
        "membership or no mapped hazard exposure, not missing data. Results do not form a "
        "Top-20 list or a repair order and do not include multi-section interactions, repair "
        "time, cost, or engineering condition."
    )
    note.font = Font(name="Arial", size=9, italic=True, color="37474F")
    note.fill = PatternFill("solid", fgColor=note_fill)
    note.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[note_row].height = 66

    widths = [44, 17, 16, 16, 18, 19, 17, 17, 17, 17, 20, 18]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:L{last_data_row}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(path)


def main() -> None:
    frame = load_analysis()
    table = summarize(frame)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        table.to_excel(
            writer, sheet_name="Network Characteristics", startrow=2, index=False
        )
    style_workbook(OUTPUT, len(table))
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Rows: {len(table)}; columns: {len(table.columns)}")
    print(table.to_string(index=False))


if __name__ == "__main__":
    main()
