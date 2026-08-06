#!/usr/bin/env python3
"""Municipal Emergency Access Reliability.

Plan: Report geographic heterogeneity in 30-minute emergency-access reliability
for every Kumamoto municipality and formal road-failure scenario.
Framework: AnaSOP Sections 5.2, 6.2-6.3, and Analytical Workflow step 5.
"""

from __future__ import annotations

import json
from pathlib import Path
from time import perf_counter

import geopandas as gpd
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from monte_carlo_emergency_routing import build_compact_emergency_network


ROOT = Path(__file__).resolve().parents[2]
PROCESSED = ROOT / "data" / "processed"
EXPERIMENT = ROOT / "data" / "exp" / "monte_carlo_length_weighted_full_1000"
CACHE = EXPERIMENT / "municipal_replicate_metrics.parquet"
OUTPUT = (
    ROOT
    / "data"
    / "results"
    / "tables"
    / "Table_municipal_emergency_access_reliability.xlsx"
)
PRIMARY_THRESHOLD = 30
PROJECTED_CRS = "EPSG:6670"
EXPECTED_SEVERITIES = (0.01, 0.03, 0.05, 0.10)
EXPECTED_REPLICATES = 1000

MUNICIPALITY_ENGLISH = {
    "あさぎり町": "Asagiri Town",
    "上天草市": "Kamiamakusa City",
    "五木村": "Itsuki Village",
    "人吉市": "Hitoyoshi City",
    "八代市": "Yatsushiro City",
    "南小国町": "Minamioguni Town",
    "南関町": "Nankan Town",
    "南阿蘇村": "Minamiaso Village",
    "合志市": "Koshi City",
    "和水町": "Nagomi Town",
    "嘉島町": "Kashima Town",
    "多良木町": "Taragi Town",
    "大津町": "Ozu Town",
    "天草市": "Amakusa City",
    "宇土市": "Uto City",
    "宇城市": "Uki City",
    "小国町": "Oguni Town",
    "山江村": "Yamae Village",
    "山都町": "Yamato Town",
    "山鹿市": "Yamaga City",
    "御船町": "Mifune Town",
    "水上村": "Mizukami Village",
    "水俣市": "Minamata City",
    "氷川町": "Hikawa Town",
    "津奈木町": "Tsunagi Town",
    "湯前町": "Yunomae Town",
    "熊本市": "Kumamoto City",
    "玉名市": "Tamana City",
    "玉東町": "Gyokuto Town",
    "球磨村": "Kuma Village",
    "産山村": "Ubuyama Village",
    "甲佐町": "Kosa Town",
    "益城町": "Mashiki Town",
    "相良村": "Sagara Village",
    "美里町": "Misato Town",
    "芦北町": "Ashikita Town",
    "苓北町": "Reihoku Town",
    "荒尾市": "Arao City",
    "菊池市": "Kikuchi City",
    "菊陽町": "Kikuyo Town",
    "西原村": "Nishihara Village",
    "錦町": "Nishiki Town",
    "長洲町": "Nagasu Town",
    "阿蘇市": "Aso City",
    "高森町": "Takamori Town",
}

TABLE_COLUMNS = [
    "Municipality Name",
    "Expected Failed Road Length (%)",
    "Total Population",
    "Population Age 65+",
    "Mean 30-Minute Timely Access Probability (%)",
    "Population-Weighted Mean Grid P90 Access Time (min)",
    "Mean Population Losing 30-Minute Access",
    "Coverage Loss Monte Carlo SE",
    "Coverage Loss P05",
    "Coverage Loss P95",
    "Mean Older Population Losing 30-Minute Access",
    "Mean Population Newly Disconnected",
]


def expected_failed_length_share(intensity: float, lengths: np.ndarray) -> float:
    probabilities = -np.expm1(-intensity * lengths)
    return float(np.dot(lengths, probabilities) / lengths.sum())


def calibrate_intensity(target: float, lengths: np.ndarray) -> float:
    low = 0.0
    high = 1.0 / float(np.median(lengths))
    while expected_failed_length_share(high, lengths) < target:
        high *= 2.0
    for _ in range(80):
        midpoint = 0.5 * (low + high)
        if expected_failed_length_share(midpoint, lengths) < target:
            low = midpoint
        else:
            high = midpoint
    return 0.5 * (low + high)


def municipality_lookup(
    features: gpd.GeoDataFrame,
    identifier: str,
    municipalities: gpd.GeoDataFrame,
) -> pd.DataFrame:
    points = features[[identifier, "Geometry"]].to_crs(PROJECTED_CRS).copy()
    administrative = municipalities[
        ["Municipality Code", "Municipality Name", "Geometry"]
    ].to_crs(PROJECTED_CRS)
    joined = gpd.sjoin(points, administrative, how="left", predicate="within")
    if joined[identifier].duplicated().any():
        raise RuntimeError(f"Ambiguous municipality assignment for {identifier}")

    missing = joined["Municipality Code"].isna()
    if missing.any():
        nearest = gpd.sjoin_nearest(
            points.loc[missing, [identifier, "Geometry"]],
            administrative,
            how="left",
            max_distance=1000.0,
            distance_col="Municipality Assignment Distance (m)",
        )
        if nearest[identifier].duplicated().any():
            raise RuntimeError(f"Ambiguous nearest municipality for {identifier}")
        nearest = nearest.set_index(identifier)
        joined = joined.set_index(identifier)
        joined.loc[
            nearest.index, ["Municipality Code", "Municipality Name"]
        ] = nearest[["Municipality Code", "Municipality Name"]]
        joined = joined.reset_index()

    if joined["Municipality Name"].isna().any():
        raise RuntimeError(f"Unassigned municipality for {identifier}")
    joined["Municipality Name"] = joined["Municipality Name"].map(
        MUNICIPALITY_ENGLISH
    )
    if joined["Municipality Name"].isna().any():
        raise RuntimeError("A municipality lacks an English label")
    return joined[[identifier, "Municipality Name"]]


def demand_dimensions() -> tuple[pd.DataFrame, pd.DataFrame, list[str]]:
    municipalities = gpd.read_parquet(
        PROCESSED / "kumamoto_administrative_areas_preprocessed.parquet",
        columns=["Municipality Code", "Municipality Name", "Geometry"],
    )
    mesh = gpd.read_parquet(
        PROCESSED / "kumamoto_population_mesh_network_access_preprocessed.parquet",
        columns=[
            "Mesh Code",
            "Analysis Unit ID",
            "Demand Node ID",
            "Total Population",
            "Geometry",
        ],
    )
    older = gpd.read_parquet(
        PROCESSED / "kumamoto_population_group_network_access_preprocessed.parquet",
        columns=[
            "Disclosure Group Code",
            "Analysis Unit ID",
            "Demand Node ID",
            "Population Age 65+",
            "Geometry",
        ],
    )
    mesh["_Demand Order"] = np.arange(len(mesh), dtype=np.int64)
    older["_Demand Order"] = np.arange(len(older), dtype=np.int64)
    mesh_lookup = municipality_lookup(mesh, "Mesh Code", municipalities)
    older_lookup = municipality_lookup(
        older, "Disclosure Group Code", municipalities
    )
    mesh = pd.DataFrame(mesh.drop(columns="Geometry")).merge(
        mesh_lookup, on="Mesh Code", how="left", validate="one_to_one"
    )
    older = pd.DataFrame(older.drop(columns="Geometry")).merge(
        older_lookup,
        on="Disclosure Group Code",
        how="left",
        validate="one_to_one",
    )
    mesh = (
        mesh.sort_values("_Demand Order", kind="stable")
        .drop(columns="_Demand Order")
        .reset_index(drop=True)
    )
    older = (
        older.sort_values("_Demand Order", kind="stable")
        .drop(columns="_Demand Order")
        .reset_index(drop=True)
    )
    municipality_names = sorted(MUNICIPALITY_ENGLISH.values())
    if len(municipality_names) != 45:
        raise RuntimeError("Expected exactly 45 English municipality names")
    if set(mesh["Municipality Name"]) != set(municipality_names):
        raise RuntimeError("Mesh demand does not cover all 45 municipalities")
    if set(older["Municipality Name"]) != set(municipality_names):
        raise RuntimeError("Older-population demand does not cover all 45 municipalities")
    return mesh, older, municipality_names


def replay_municipal_metrics() -> pd.DataFrame:
    summary = json.loads((EXPERIMENT / "experiment_summary.json").read_text())
    targets = np.asarray(
        summary["design"]["expected_failed_road_length_shares"], dtype=float
    )
    replicates = int(summary["design"]["replicates_per_severity"])
    base_seed = int(summary["design"]["base_seed"])
    if not np.allclose(targets, EXPECTED_SEVERITIES) or replicates != EXPECTED_REPLICATES:
        raise RuntimeError("The formal experiment design changed")

    mesh, older, municipality_names = demand_dimensions()
    municipality_index = {name: index for index, name in enumerate(municipality_names)}

    network = build_compact_emergency_network(
        PROCESSED, include_population_groups=True
    )
    mesh_mask = network.demand_type == "population_mesh"
    older_mask = network.demand_type == "older_population_group"
    mesh_count = int(mesh_mask.sum())
    older_count = int(older_mask.sum())
    if mesh_count != len(mesh) or older_count != len(older):
        raise RuntimeError("Demand dimensions do not match the routing network")
    mesh_sequence = (
        mesh["Demand Node ID"].astype("string").fillna("").astype(str).tolist()
    )
    older_sequence = (
        older["Demand Node ID"].astype("string").fillna("").astype(str).tolist()
    )
    mesh_ids = [str(value) for value in network.demand_ids[:mesh_count]]
    older_ids = [str(value) for value in network.demand_ids[mesh_count:]]
    if mesh_sequence != mesh_ids:
        first = next(
            index
            for index, (observed, expected) in enumerate(
                zip(mesh_sequence, mesh_ids, strict=True)
            )
            if observed != expected
        )
        raise RuntimeError(
            f"Mesh demand sequence differs at row {first}: "
            f"{mesh_sequence[first]} != {mesh_ids[first]}"
        )
    if older_sequence != older_ids:
        first = next(
            index
            for index, (observed, expected) in enumerate(
                zip(older_sequence, older_ids, strict=True)
            )
            if observed != expected
        )
        raise RuntimeError(
            f"Older-population demand sequence differs at row {first}: "
            f"{older_sequence[first]} != {older_ids[first]}"
        )
    mesh_municipality = mesh["Municipality Name"].map(municipality_index).to_numpy(int)
    older_municipality = older["Municipality Name"].map(municipality_index).to_numpy(int)
    baseline = network.route(assign_hospitals=False)

    sections = pd.read_parquet(
        PROCESSED / "kumamoto_road_sections_preprocessed.parquet",
        columns=["Road Section ID", "Road Section Length (m)"],
    )
    sections["Road Section ID"] = sections["Road Section ID"].astype(str)
    sections = sections.set_index("Road Section ID").loc[
        network.road_section_ids
    ]
    lengths = sections["Road Section Length (m)"].to_numpy(float)
    intensities = np.asarray(
        [calibrate_intensity(target, lengths) for target in targets]
    )
    probabilities = np.vstack(
        [-np.expm1(-intensity * lengths) for intensity in intensities]
    )

    population = network.demand_population[:mesh_count]
    older_population = network.demand_older_population[mesh_count:]
    municipality_count = len(municipality_names)
    total_population = np.bincount(
        mesh_municipality, weights=population, minlength=municipality_count
    )
    population_age_65 = np.bincount(
        older_municipality,
        weights=older_population,
        minlength=municipality_count,
    )
    baseline_mesh_time = baseline.total_time[:mesh_count]
    baseline_older_time = baseline.total_time[mesh_count:]
    baseline_mesh_timely = baseline_mesh_time <= PRIMARY_THRESHOLD
    baseline_older_timely = baseline_older_time <= PRIMARY_THRESHOLD
    baseline_mesh_finite = np.isfinite(baseline_mesh_time)

    row_count = len(targets) * replicates * municipality_count
    scenario_values = np.empty(row_count, dtype=np.float64)
    replicate_values = np.empty(row_count, dtype=np.int16)
    municipality_values = np.empty(row_count, dtype=object)
    coverage_values = np.empty(row_count, dtype=np.float64)
    loss_values = np.empty(row_count, dtype=np.float64)
    older_loss_values = np.empty(row_count, dtype=np.float64)
    disconnected_values = np.empty(row_count, dtype=np.float64)
    cursor = 0
    started = perf_counter()

    for replicate in range(1, replicates + 1):
        rng = np.random.default_rng(np.random.SeedSequence([base_seed, replicate]))
        uniform_score = rng.random(network.road_section_count)
        previous_failed = np.zeros(network.road_section_count, dtype=bool)
        for severity_index, target in enumerate(targets):
            failed = uniform_score < probabilities[severity_index]
            if not np.all(failed[previous_failed]):
                raise RuntimeError("Nested failure assignment failed")
            state = network.route(failed, assign_hospitals=False)
            mesh_time = state.total_time[:mesh_count]
            older_time = state.total_time[mesh_count:]
            mesh_timely = mesh_time <= PRIMARY_THRESHOLD
            older_timely = older_time <= PRIMARY_THRESHOLD
            mesh_lost = baseline_mesh_timely & ~mesh_timely
            older_lost = baseline_older_timely & ~older_timely
            newly_disconnected = baseline_mesh_finite & ~np.isfinite(mesh_time)

            covered_by_municipality = np.bincount(
                mesh_municipality,
                weights=population * mesh_timely,
                minlength=municipality_count,
            )
            loss_by_municipality = np.bincount(
                mesh_municipality,
                weights=population * mesh_lost,
                minlength=municipality_count,
            )
            older_loss_by_municipality = np.bincount(
                older_municipality,
                weights=older_population * older_lost,
                minlength=municipality_count,
            )
            disconnected_by_municipality = np.bincount(
                mesh_municipality,
                weights=population * newly_disconnected,
                minlength=municipality_count,
            )

            stop = cursor + municipality_count
            scenario_values[cursor:stop] = target
            replicate_values[cursor:stop] = replicate
            municipality_values[cursor:stop] = municipality_names
            coverage_values[cursor:stop] = covered_by_municipality
            loss_values[cursor:stop] = loss_by_municipality
            older_loss_values[cursor:stop] = older_loss_by_municipality
            disconnected_values[cursor:stop] = disconnected_by_municipality
            cursor = stop
            previous_failed = failed

        if replicate == 1 or replicate % 25 == 0 or replicate == replicates:
            elapsed = perf_counter() - started
            print(
                f"Municipal replay: {replicate}/{replicates} replicates "
                f"({len(targets) * replicate:,} states), {elapsed:.1f}s elapsed",
                flush=True,
            )

    if cursor != row_count:
        raise RuntimeError("Municipal replay output length mismatch")
    output = pd.DataFrame(
        {
            "Municipality Name": municipality_values,
            "Expected Failed Road Length Share": scenario_values,
            "Simulation Replicate": replicate_values,
            "Population within 30 Minutes": coverage_values,
            "Population Losing Baseline 30-Minute Access": loss_values,
            "Older Population Losing Baseline 30-Minute Access": older_loss_values,
            "Population Newly Disconnected": disconnected_values,
        }
    )
    population_dimension = pd.DataFrame(
        {
            "Municipality Name": municipality_names,
            "Total Population": total_population,
            "Population Age 65+": population_age_65,
        }
    )
    output = output.merge(
        population_dimension,
        on="Municipality Name",
        how="left",
        validate="many_to_one",
    )
    validate_against_formal_metrics(output)
    temporary = CACHE.with_suffix(".tmp.parquet")
    output.to_parquet(temporary, index=False)
    temporary.replace(CACHE)
    print(f"Saved replay cache: {CACHE.relative_to(ROOT)}", flush=True)
    return output


def validate_against_formal_metrics(municipal: pd.DataFrame) -> None:
    formal = pd.read_parquet(EXPERIMENT / "replicate_metrics.parquet")
    keys = ["Expected Failed Road Length Share", "Simulation Replicate"]
    aggregate = (
        municipal.groupby(keys, as_index=False)
        .agg(
            **{
                "Population within 30 Minutes": (
                    "Population within 30 Minutes",
                    "sum",
                ),
                "Population Losing Baseline 30-Minute Access": (
                    "Population Losing Baseline 30-Minute Access",
                    "sum",
                ),
                "Older Population Losing Baseline 30-Minute Access": (
                    "Older Population Losing Baseline 30-Minute Access",
                    "sum",
                ),
                "Population Newly Disconnected": (
                    "Population Newly Disconnected",
                    "sum",
                ),
            }
        )
        .sort_values(keys)
        .reset_index(drop=True)
    )
    formal = formal.sort_values(keys).reset_index(drop=True)
    for column in [
        "Population within 30 Minutes",
        "Population Losing Baseline 30-Minute Access",
        "Older Population Losing Baseline 30-Minute Access",
        "Population Newly Disconnected",
    ]:
        if not np.allclose(aggregate[column], formal[column], atol=1e-6):
            difference = float(np.max(np.abs(aggregate[column] - formal[column])))
            raise RuntimeError(
                f"Municipal replay does not reproduce {column}; max difference {difference}"
            )


def load_or_generate_cache() -> pd.DataFrame:
    if CACHE.exists():
        municipal = pd.read_parquet(CACHE)
        required = {
            "Municipality Name",
            "Expected Failed Road Length Share",
            "Simulation Replicate",
            "Population within 30 Minutes",
            "Population Losing Baseline 30-Minute Access",
            "Older Population Losing Baseline 30-Minute Access",
            "Population Newly Disconnected",
            "Total Population",
            "Population Age 65+",
        }
        if required.issubset(municipal.columns) and len(municipal) == 180000:
            municipal["Expected Failed Road Length Share"] = (
                municipal["Expected Failed Road Length Share"].astype(float).round(8)
            )
            validate_against_formal_metrics(municipal)
            print(f"Using replay cache: {CACHE.relative_to(ROOT)}", flush=True)
            return municipal
        raise RuntimeError("Existing municipal replay cache is incompatible")
    return replay_municipal_metrics()


def grid_p90_summary() -> pd.DataFrame:
    municipalities = gpd.read_parquet(
        PROCESSED / "kumamoto_administrative_areas_preprocessed.parquet",
        columns=["Municipality Code", "Municipality Name", "Geometry"],
    )
    mesh = gpd.read_parquet(
        PROCESSED / "kumamoto_population_mesh_network_access_preprocessed.parquet",
        columns=["Mesh Code", "Geometry"],
    )
    lookup = municipality_lookup(mesh, "Mesh Code", municipalities)
    grid = pd.read_parquet(
        EXPERIMENT / "grid_reliability.parquet",
        columns=[
            "Mesh Code",
            "Total Population",
            "Expected Failed Road Length Share",
            "P90 Emergency Access Time",
            "P90 Unreachable",
        ],
    ).merge(lookup, on="Mesh Code", how="left", validate="many_to_one")
    grid["Expected Failed Road Length Share"] = (
        grid["Expected Failed Road Length Share"].astype(float).round(8)
    )
    valid = ~grid["P90 Unreachable"] & grid["P90 Emergency Access Time"].notna()
    grid["P90 Weighted Time"] = (
        grid["Total Population"] * grid["P90 Emergency Access Time"]
    ).where(valid, 0.0)
    grid["P90 Weight"] = grid["Total Population"].where(valid, 0.0)
    summary = (
        grid.groupby(
            ["Municipality Name", "Expected Failed Road Length Share"],
            as_index=False,
        )
        .agg(
            **{
                "P90 Weighted Time": ("P90 Weighted Time", "sum"),
                "P90 Weight": ("P90 Weight", "sum"),
            }
        )
    )
    summary["Population-Weighted Mean Grid P90 Access Time (min)"] = (
        summary["P90 Weighted Time"] / summary["P90 Weight"].replace(0, np.nan)
    )
    return summary[
        [
            "Municipality Name",
            "Expected Failed Road Length Share",
            "Population-Weighted Mean Grid P90 Access Time (min)",
        ]
    ]


def build_table(municipal: pd.DataFrame) -> pd.DataFrame:
    municipal = municipal.copy()
    municipal["Expected Failed Road Length Share"] = (
        municipal["Expected Failed Road Length Share"].astype(float).round(8)
    )
    keys = ["Municipality Name", "Expected Failed Road Length Share"]
    grouped = municipal.groupby(keys, sort=True)
    table = grouped.agg(
        **{
            "Total Population": ("Total Population", "first"),
            "Population Age 65+": ("Population Age 65+", "first"),
            "Mean Population with 30-Minute Access": (
                "Population within 30 Minutes",
                "mean",
            ),
            "Mean Population Losing 30-Minute Access": (
                "Population Losing Baseline 30-Minute Access",
                "mean",
            ),
            "Coverage Loss Standard Deviation": (
                "Population Losing Baseline 30-Minute Access",
                "std",
            ),
            "Coverage Loss P05": (
                "Population Losing Baseline 30-Minute Access",
                lambda series: series.quantile(0.05),
            ),
            "Coverage Loss P95": (
                "Population Losing Baseline 30-Minute Access",
                lambda series: series.quantile(0.95),
            ),
            "Mean Older Population Losing 30-Minute Access": (
                "Older Population Losing Baseline 30-Minute Access",
                "mean",
            ),
            "Mean Population Newly Disconnected": (
                "Population Newly Disconnected",
                "mean",
            ),
            "Simulation Replicates": ("Simulation Replicate", "nunique"),
        }
    ).reset_index()
    table["Coverage Loss Monte Carlo SE"] = (
        table["Coverage Loss Standard Deviation"]
        / np.sqrt(table["Simulation Replicates"])
    )
    table["Mean 30-Minute Timely Access Probability (%)"] = 100.0 * (
        table["Mean Population with 30-Minute Access"]
        / table["Total Population"].replace(0, np.nan)
    )
    table = table.merge(
        grid_p90_summary(), on=keys, how="left", validate="one_to_one"
    )
    p90_column = "Population-Weighted Mean Grid P90 Access Time (min)"
    missing_p90 = table[p90_column].isna()
    expected_missing = (
        table["Municipality Name"].eq("Itsuki Village")
        & np.isclose(table["Expected Failed Road Length Share"], 0.10)
    )
    if not missing_p90.equals(expected_missing):
        raise RuntimeError("Unexpected pattern of unavailable municipal P90 values")
    table[p90_column] = table[p90_column].astype(object).where(~missing_p90, "None")
    table["Expected Failed Road Length (%)"] = (
        100.0 * table["Expected Failed Road Length Share"]
    )
    table = table.sort_values(
        ["Municipality Name", "Expected Failed Road Length Share"],
        kind="stable",
    ).reset_index(drop=True)
    table = table[TABLE_COLUMNS]
    if table.shape != (180, 12):
        raise RuntimeError(f"Unexpected municipal table shape: {table.shape}")
    if table.isna().any().any():
        raise RuntimeError("Municipal table contains unexpected missing values")
    return table


def style_workbook(path: Path, data_rows: int) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook["Municipal Reliability"]
    navy = "1F4E78"
    main_fill = "F4F8FB"
    main_alt_fill = "E8F1F7"
    stress_fill = "FCE4D6"
    note_fill = "EAF3F8"
    white = "FFFFFF"
    thin_grey = Side(style="thin", color="B7C3CA")
    medium_grey = Side(style="medium", color="7F8C8D")

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
    title = sheet.cell(1, 1)
    title.value = "Municipal Emergency Access Reliability"
    title.font = Font(name="Arial", size=14, bold=True, color=white)
    title.fill = PatternFill("solid", fgColor=navy)
    title.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 24

    header_row = 3
    for cell in sheet[header_row]:
        cell.font = Font(name="Arial", size=8.5, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=thin_grey,
            bottom=thin_grey,
            left=thin_grey,
            right=thin_grey,
        )
    sheet.row_dimensions[header_row].height = 72

    first_data_row = 4
    last_data_row = first_data_row + data_rows - 1
    municipality_group = -1
    previous_municipality: str | None = None
    for row in range(first_data_row, last_data_row + 1):
        municipality = str(sheet.cell(row, 1).value)
        severity = float(sheet.cell(row, 2).value)
        starts_group = municipality != previous_municipality
        if starts_group:
            municipality_group += 1
        if np.isclose(severity, 10.0):
            fill = PatternFill("solid", fgColor=stress_fill)
        else:
            fill = PatternFill(
                "solid", fgColor=main_fill if municipality_group % 2 == 0 else main_alt_fill
            )
        top_side = medium_grey if starts_group and row > first_data_row else thin_grey
        for column in range(1, 13):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Arial", size=8.5)
            cell.fill = fill
            cell.border = Border(
                top=top_side,
                bottom=thin_grey,
                left=thin_grey,
                right=thin_grey,
            )
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "right",
                vertical="center",
            )
        sheet.cell(row, 2).number_format = '0"%"'
        for column in (3, 4, 7, 9, 10, 11, 12):
            sheet.cell(row, column).number_format = "#,##0"
        sheet.cell(row, 5).number_format = '0.0"%"'
        sheet.cell(row, 6).number_format = "0.0"
        sheet.cell(row, 8).number_format = "#,##0.0"
        previous_municipality = municipality

    note_row = last_data_row + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=12)
    note = sheet.cell(note_row, 1)
    note.value = (
        "Note: Results use the 30-minute threshold and 1,000 paired Monte Carlo "
        "replicates per scenario. The 1%, 3%, and 5% states are main scenarios; 10% is "
        "an extreme stress scenario. P05, P95, and Monte Carlo SE summarize municipal "
        "population loss across replicates. Population-weighted mean grid P90 time "
        "averages grid-level P90 values among grids whose P90 remains finite; older "
        "population uses valid disclosure-group support. Kumamoto City's five wards are "
        "combined. None means that no grid in the municipality has an estimable finite "
        "P90; this occurs for Itsuki Village in the 10% stress scenario because every "
        "grid is P90-unreachable."
    )
    note.font = Font(name="Arial", size=9, italic=True, color="37474F")
    note.fill = PatternFill("solid", fgColor=note_fill)
    note.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[note_row].height = 54

    widths = [24, 17, 16, 17, 21, 24, 22, 18, 17, 17, 24, 22]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:L{last_data_row}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(path)


def main() -> None:
    municipal = load_or_generate_cache()
    table = build_table(municipal)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        table.to_excel(
            writer, sheet_name="Municipal Reliability", startrow=2, index=False
        )
    style_workbook(OUTPUT, len(table))
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Rows: {len(table)}; columns: {len(table.columns)}")
    print(table.head(12).to_string(index=False))


if __name__ == "__main__":
    main()
