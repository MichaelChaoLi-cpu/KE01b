#!/usr/bin/env python3
"""Hospital Service Reliability.

Plan: Report hospital assignment and service-population stability under the
three main road-failure scenarios and the 10% stress scenario.
Framework: AnaSOP Sections 5.2, 6.4, and Analytical Workflow step 6.
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
PROCESSED = ROOT / "data" / "processed"
EXPERIMENT = ROOT / "data" / "exp" / "monte_carlo_length_weighted_full_1000"
INPUT = EXPERIMENT / "hospital_service_reliability.parquet"
OUTPUT = (
    ROOT
    / "data"
    / "results"
    / "tables"
    / "Table_hospital_service_reliability.xlsx"
)

JAPANESE_CHARACTER_PATTERN = re.compile(
    r"[\u3040-\u30ff\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]"
)
HOSPITAL_ENGLISH_OVERRIDES = {
    "HOSPITAL-0000006": "Mimori Cardiovascular and Respiratory Hospital",
    "HOSPITAL-0000010": "Hori Hospital",
    "HOSPITAL-0000012": "Public Tamana Central Hospital",
    "HOSPITAL-0000013": "Terao Hospital",
    "HOSPITAL-0000018": "Minamata City General Medical Center",
    "HOSPITAL-0000020": "Uki Municipal Hospital",
    "HOSPITAL-0000025": "Amakusa Daiichi Hospital",
    "HOSPITAL-0000026": "Uki General Hospital",
    "HOSPITAL-0000031": "Yamaga City Medical Center",
    "HOSPITAL-0000032": "Okabe Hospital",
    "HOSPITAL-0000048": "Saiseikai Misumi Hospital",
    "HOSPITAL-0000052": "Kumamoto Chuo Hospital",
    "HOSPITAL-0000060": "JCHO Hitoyoshi Medical Center",
    "HOSPITAL-0000061": "Kumamoto Rosai Hospital",
    "HOSPITAL-0000063": "JCHO Amakusa Chuo General Hospital",
    "HOSPITAL-0000064": "Tamana Regional Health Medical Center",
    "HOSPITAL-0000065": "Kuma District Public Taragi Hospital",
    "HOSPITAL-0000067": "Arao Municipal Hospital",
    "HOSPITAL-0000068": "Kikuchi Chuo Hospital",
}
TABLE_COLUMNS = [
    "Hospital ID",
    "Hospital Name",
    "Hospital Role",
    "Expected Failed Road Length (%)",
    "Baseline Hospital Catchment Population",
    "Mean Hospital Catchment Population",
    "Population-Weighted Hospital Assignment Probability (%)",
    "Hospital Demand Change",
    "Hospital Catchment P05",
    "Hospital Catchment P50",
    "Hospital Catchment P95",
    "Zero Catchment Frequency (%)",
]


def normalized_name(value: object) -> str:
    if pd.isna(value):
        return ""
    normalized = unicodedata.normalize("NFKC", str(value).strip()).casefold()
    return re.sub(r"[\s・·.,，。、()（）\-]", "", normalized)


def hospital_english_names(hospitals: pd.DataFrame) -> dict[str, str]:
    reference = pd.read_parquet(
        PROCESSED / "kumamoto_hospitals_2026_preprocessed.parquet",
        columns=["Hospital Name", "Hospital Name Romanized"],
    ).dropna(subset=["Hospital Name Romanized"])
    reference["_Normalized Name"] = reference["Hospital Name"].map(normalized_name)
    reference = reference.loc[reference["_Normalized Name"].ne("")].copy()
    exact = dict(
        zip(
            reference["_Normalized Name"],
            reference["Hospital Name Romanized"].astype(str).str.strip(),
            strict=True,
        )
    )
    reference_pairs = list(exact.items())

    labels: dict[str, str] = {}
    unique = hospitals[["Hospital Node ID", "Facility Name"]].drop_duplicates()
    for hospital_node_id, facility_name in unique.itertuples(index=False, name=None):
        hospital_node_id = str(hospital_node_id)
        if hospital_node_id in HOSPITAL_ENGLISH_OVERRIDES:
            label = HOSPITAL_ENGLISH_OVERRIDES[hospital_node_id]
        else:
            normalized_facility = normalized_name(facility_name)
            label = exact.get(normalized_facility)
            if label is None:
                candidates = [
                    (len(reference_name), reference_romanized)
                    for reference_name, reference_romanized in reference_pairs
                    if len(reference_name) >= 5
                    and (
                        reference_name in normalized_facility
                        or normalized_facility in reference_name
                    )
                ]
                if candidates:
                    label = sorted(
                        candidates, key=lambda item: (-item[0], item[1])
                    )[0][1]
            if label is None:
                label = f"Hospital {hospital_node_id.replace('HOSPITAL-', '')}"
        label = str(label).strip()
        if JAPANESE_CHARACTER_PATTERN.search(label):
            raise RuntimeError(
                f"Japanese characters remain in {hospital_node_id}: {label}"
            )
        labels[hospital_node_id] = label

    if len(labels) != 75:
        raise RuntimeError(f"Expected 75 English hospital labels; got {len(labels)}")
    if len(set(labels.values())) != 75:
        raise RuntimeError("English hospital labels are not unique")
    return labels


def hospital_role(frame: pd.DataFrame) -> pd.Series:
    emergency = frame["Emergency Designation"].eq("Designated")
    core = frame["Disaster Base Designation"].eq("Core Disaster Base")
    regional = frame["Disaster Base Designation"].eq("Regional Disaster Base")
    role = pd.Series("Emergency Hospital", index=frame.index, dtype="string")
    role.loc[regional & emergency] = "Regional Disaster Base Emergency Hospital"
    role.loc[regional & ~emergency] = "Regional Disaster Base Hospital"
    role.loc[core] = "Core Disaster Base Emergency Hospital"
    if role.isna().any():
        raise RuntimeError("A hospital role could not be constructed")
    return role


def validate_source(frame: pd.DataFrame) -> None:
    if frame.shape[0] != 300:
        raise RuntimeError(f"Expected 300 hospital-scenario rows; got {len(frame)}")
    if frame["Hospital Node ID"].nunique() != 75:
        raise RuntimeError("Expected 75 eligible hospitals")
    if not frame["Simulation Replicates"].eq(1000).all():
        raise RuntimeError("Hospital results do not all use 1,000 replicates")
    expected_severities = {0.01, 0.03, 0.05, 0.10}
    observed_severities = set(
        frame["Expected Failed Road Length Share"].astype(float).round(8)
    )
    if observed_severities != expected_severities:
        raise RuntimeError("Hospital results do not contain the four formal severities")
    hospital_counts = frame.groupby("Hospital Node ID").size()
    if not hospital_counts.eq(4).all():
        raise RuntimeError("Each hospital must have four severity rows")
    baseline_counts = frame.groupby("Hospital Node ID")[
        "Baseline Hospital Catchment Population"
    ].nunique()
    if not baseline_counts.eq(1).all():
        raise RuntimeError("Baseline hospital catchment changes across scenarios")
    if not (
        frame["Hospital Catchment P05"]
        <= frame["Hospital Catchment P50"]
    ).all() or not (
        frame["Hospital Catchment P50"]
        <= frame["Hospital Catchment P95"]
    ).all():
        raise RuntimeError("Hospital catchment quantiles are not ordered")
    if not frame["Zero Catchment Frequency"].between(0, 1).all():
        raise RuntimeError("Zero catchment frequency is outside [0, 1]")

    total_population = float(
        pd.read_parquet(
            PROCESSED / "kumamoto_population_mesh_network_access_preprocessed.parquet",
            columns=["Total Population"],
        )["Total Population"].sum()
    )
    expected_change = (
        frame["Mean Hospital Catchment Population"]
        - frame["Baseline Hospital Catchment Population"]
    )
    if not np.allclose(expected_change, frame["Hospital Demand Change"], atol=1e-9):
        raise RuntimeError("Hospital demand change identity failed")
    expected_probability = (
        frame["Mean Hospital Catchment Population"] / total_population
    )
    if not np.allclose(
        expected_probability,
        frame["Population-Weighted Hospital Assignment Probability"],
        atol=1e-12,
    ):
        raise RuntimeError("Hospital assignment probability identity failed")
    if (frame["Hospital Catchment Standard Deviation"] < 0).any():
        raise RuntimeError("Hospital catchment standard deviation is negative")
    catchment_totals = (
        frame.groupby("Expected Failed Road Length Share")[
            "Mean Hospital Catchment Population"
        ]
        .sum()
        .sort_index()
    )
    if (np.diff(catchment_totals.to_numpy(float)) > 1e-6).any():
        raise RuntimeError("Total hospital-assigned population increases with severity")
    baseline_sum = (
        frame.drop_duplicates("Hospital Node ID")[
            "Baseline Hospital Catchment Population"
        ].sum()
    )
    if baseline_sum > total_population + 1e-6:
        raise RuntimeError("Baseline hospital catchments exceed total population")
    if (catchment_totals > baseline_sum + 1e-6).any():
        raise RuntimeError("A disrupted catchment total exceeds the baseline total")


def build_table(frame: pd.DataFrame) -> pd.DataFrame:
    labels = hospital_english_names(frame)
    table = pd.DataFrame(
        {
            "Hospital ID": frame["Hospital Node ID"].astype(str),
            "Hospital Name": frame["Hospital Node ID"].astype(str).map(labels),
            "Hospital Role": hospital_role(frame),
            "Expected Failed Road Length (%)": 100.0
            * frame["Expected Failed Road Length Share"].astype(float),
            "Baseline Hospital Catchment Population": frame[
                "Baseline Hospital Catchment Population"
            ].astype(float),
            "Mean Hospital Catchment Population": frame[
                "Mean Hospital Catchment Population"
            ].astype(float),
            "Population-Weighted Hospital Assignment Probability (%)": 100.0
            * frame[
                "Population-Weighted Hospital Assignment Probability"
            ].astype(float),
            "Hospital Demand Change": frame["Hospital Demand Change"].astype(float),
            "Hospital Catchment P05": frame["Hospital Catchment P05"].astype(float),
            "Hospital Catchment P50": frame["Hospital Catchment P50"].astype(float),
            "Hospital Catchment P95": frame["Hospital Catchment P95"].astype(float),
            "Zero Catchment Frequency (%)": 100.0
            * frame["Zero Catchment Frequency"].astype(float),
        }
    )
    table = table.sort_values(
        ["Hospital Name", "Expected Failed Road Length (%)"], kind="stable"
    ).reset_index(drop=True)
    table = table[TABLE_COLUMNS]
    if table.shape != (300, 12):
        raise RuntimeError(f"Unexpected hospital table shape: {table.shape}")
    if table.isna().any().any():
        raise RuntimeError("Hospital table contains missing values")
    if any(
        JAPANESE_CHARACTER_PATTERN.search(str(value))
        for column in ["Hospital ID", "Hospital Name", "Hospital Role"]
        for value in table[column]
    ):
        raise RuntimeError("Japanese characters remain in the hospital table")
    return table


def style_workbook(path: Path, data_rows: int) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook["Hospital Reliability"]
    navy = "1F4E78"
    main_fill = "F4F8FB"
    main_alt_fill = "E8F1F7"
    stress_fill = "FCE4D6"
    note_fill = "EAF3F8"
    white = "FFFFFF"
    thin_grey = Side(style="thin", color="B7C3CA")
    medium_grey = Side(style="medium", color="7F8C8D")

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
    title = sheet.cell(1, 1)
    title.value = "Hospital Service Reliability"
    title.font = Font(name="Arial", size=14, bold=True, color=white)
    title.fill = PatternFill("solid", fgColor=navy)
    title.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 24

    header_row = 3
    for cell in sheet[header_row]:
        cell.font = Font(name="Arial", size=8.5, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=thin_grey,
            bottom=thin_grey,
            left=thin_grey,
            right=thin_grey,
        )
    sheet.row_dimensions[header_row].height = 72

    first_data_row = 4
    last_data_row = first_data_row + data_rows - 1
    hospital_group = -1
    previous_hospital: str | None = None
    for row in range(first_data_row, last_data_row + 1):
        hospital_name = str(sheet.cell(row, 2).value)
        severity = float(sheet.cell(row, 4).value)
        starts_group = hospital_name != previous_hospital
        if starts_group:
            hospital_group += 1
        if np.isclose(severity, 10.0):
            fill = PatternFill("solid", fgColor=stress_fill)
        else:
            fill = PatternFill(
                "solid", fgColor=main_fill if hospital_group % 2 == 0 else main_alt_fill
            )
        top_side = medium_grey if starts_group and row > first_data_row else thin_grey
        for column in range(1, 13):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Arial", size=8.5)
            cell.fill = fill
            cell.border = Border(
                top=top_side,
                bottom=thin_grey,
                left=thin_grey,
                right=thin_grey,
            )
            cell.alignment = Alignment(
                horizontal="left" if column <= 3 else "right",
                vertical="center",
                wrap_text=column in (2, 3),
            )
        sheet.cell(row, 4).number_format = '0"%"'
        for column in (5, 6, 8, 9, 10, 11):
            sheet.cell(row, column).number_format = "#,##0"
        sheet.cell(row, 7).number_format = '0.000"%"'
        sheet.cell(row, 12).number_format = '0.0"%"'
        previous_hospital = hospital_name

    note_row = last_data_row + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=12)
    note = sheet.cell(note_row, 1)
    note.value = (
        "Note: Results use 1,000 paired Monte Carlo replicates per scenario. The 1%, "
        "3%, and 5% states are main scenarios; 10% is an extreme stress scenario. "
        "Catchment population is the population assigned to the nearest eligible hospital "
        "in each rerouted network state. Assignment probability is mean catchment divided "
        "by total prefectural population. Zero catchment frequency is the share of "
        "replicates with no assigned population; it does not establish hospital building "
        "failure, clinical closure, or zero treatment capacity. Hospital role and capacity "
        "do not weight destination choice in the primary specification. Names use English "
        "romanization or reviewed English labels."
    )
    note.font = Font(name="Arial", size=9, italic=True, color="37474F")
    note.fill = PatternFill("solid", fgColor=note_fill)
    note.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[note_row].height = 64

    widths = [20, 35, 35, 17, 20, 20, 22, 18, 17, 17, 17, 18]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "D4"
    sheet.auto_filter.ref = f"A3:L{last_data_row}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(path)


def main() -> None:
    frame = pd.read_parquet(INPUT)
    validate_source(frame)
    table = build_table(frame)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        table.to_excel(
            writer, sheet_name="Hospital Reliability", startrow=2, index=False
        )
    style_workbook(OUTPUT, len(table))
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Rows: {len(table)}; columns: {len(table.columns)}")
    print(table.head(12).to_string(index=False))


if __name__ == "__main__":
    main()
