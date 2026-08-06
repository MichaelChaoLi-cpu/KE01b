#!/usr/bin/env python3
"""Monte Carlo Convergence Summary.

Plan: Report stability of primary 30-minute accessibility outcomes across five
replicate checkpoints and four formal road-failure severities.
Framework: AnaSOP Sections 6.6-6.7 and Analytical Workflow step 8.
"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
EXPERIMENT = ROOT / "data" / "exp" / "monte_carlo_length_weighted_full_1000"
INPUT = EXPERIMENT / "convergence_summary.parquet"
OUTPUT = (
    ROOT
    / "data"
    / "results"
    / "tables"
    / "Table_monte_carlo_convergence_summary.xlsx"
)
CHECKPOINTS = (100, 250, 500, 750, 1000)
SEVERITIES = (0.01, 0.03, 0.05, 0.10)
COVERAGE_MEAN_CHANGE_SHARE_MAXIMUM = 0.005
GRID_PROBABILITY_P95_ABSOLUTE_CHANGE_MAXIMUM = 0.05
HOSPITAL_CATCHMENT_P95_CHANGE_SHARE_MAXIMUM = 0.005
TABLE_COLUMNS = [
    "Expected Failed Road Length (%)",
    "Simulation Replicate Checkpoint",
    "Mean Population within 30 Minutes",
    "Population Coverage Monte Carlo SE",
    "Population Coverage P05",
    "Population Coverage P95",
    "Coverage Mean Change (% of Total Population)",
    "Grid Probability P95 Change (percentage points)",
    "Hospital Catchment P95 Change (% of Total Population)",
    "Monte Carlo Convergence Status",
]


def validate_source(frame: pd.DataFrame, total_population: float) -> None:
    if frame.shape != (20, 12):
        raise RuntimeError(f"Unexpected convergence source shape: {frame.shape}")
    observed_checkpoints = set(frame["Simulation Replicate Checkpoint"].astype(int))
    observed_severities = set(
        frame["Expected Failed Road Length Share"].astype(float).round(8)
    )
    if observed_checkpoints != set(CHECKPOINTS):
        raise RuntimeError("Convergence checkpoints changed")
    if observed_severities != set(SEVERITIES):
        raise RuntimeError("Convergence severities changed")
    if not frame.groupby("Expected Failed Road Length Share").size().eq(5).all():
        raise RuntimeError("Each severity must contain five checkpoints")
    reference = frame["Simulation Replicate Checkpoint"].eq(100)
    change_columns = [
        "Coverage Mean Change Share",
        "Grid Probability P95 Absolute Change",
        "Hospital Catchment P95 Absolute Change",
    ]
    if not frame.loc[reference, change_columns].isna().all().all():
        raise RuntimeError("Reference checkpoint contains change statistics")
    if frame.loc[~reference, change_columns].isna().any().any():
        raise RuntimeError("A non-reference checkpoint lacks change statistics")
    if not frame.loc[reference, "Monte Carlo Convergence Status"].eq(
        "Reference checkpoint"
    ).all():
        raise RuntimeError("Reference checkpoint status changed")

    nonreference = frame.loc[~reference].copy()
    expected_stable = (
        nonreference["Coverage Mean Change Share"].le(
            COVERAGE_MEAN_CHANGE_SHARE_MAXIMUM
        )
        & nonreference["Grid Probability P95 Absolute Change"].le(
            GRID_PROBABILITY_P95_ABSOLUTE_CHANGE_MAXIMUM
        )
        & (
            nonreference["Hospital Catchment P95 Absolute Change"]
            / total_population
        ).le(HOSPITAL_CATCHMENT_P95_CHANGE_SHARE_MAXIMUM)
    )
    observed_stable = nonreference["Monte Carlo Convergence Status"].eq("Stable")
    if not expected_stable.equals(observed_stable):
        raise RuntimeError("Stored convergence status does not match declared thresholds")
    final = frame["Simulation Replicate Checkpoint"].eq(1000)
    if not frame.loc[final, "Monte Carlo Convergence Status"].eq("Stable").all():
        raise RuntimeError("A final severity result is not stable")
    for _, group in frame.groupby("Expected Failed Road Length Share"):
        group = group.sort_values("Simulation Replicate Checkpoint")
        if not (
            group["Population Coverage Standard Error"].iloc[-1]
            < group["Population Coverage Standard Error"].iloc[0]
        ):
            raise RuntimeError("Final coverage Monte Carlo SE did not decrease")


def build_table(frame: pd.DataFrame, total_population: float) -> pd.DataFrame:
    table = pd.DataFrame(
        {
            "Expected Failed Road Length (%)": 100.0
            * frame["Expected Failed Road Length Share"].astype(float),
            "Simulation Replicate Checkpoint": frame[
                "Simulation Replicate Checkpoint"
            ].astype(int),
            "Mean Population within 30 Minutes": frame[
                "Mean Population within 30 Minutes"
            ].astype(float),
            "Population Coverage Monte Carlo SE": frame[
                "Population Coverage Standard Error"
            ].astype(float),
            "Population Coverage P05": frame["Population Coverage P05"].astype(
                float
            ),
            "Population Coverage P95": frame["Population Coverage P95"].astype(
                float
            ),
            "Coverage Mean Change (% of Total Population)": 100.0
            * frame["Coverage Mean Change Share"].astype(float),
            "Grid Probability P95 Change (percentage points)": 100.0
            * frame["Grid Probability P95 Absolute Change"].astype(float),
            "Hospital Catchment P95 Change (% of Total Population)": 100.0
            * frame["Hospital Catchment P95 Absolute Change"].astype(float)
            / total_population,
            "Monte Carlo Convergence Status": frame[
                "Monte Carlo Convergence Status"
            ].astype(str),
        }
    )
    table = table.sort_values(
        [
            "Expected Failed Road Length (%)",
            "Simulation Replicate Checkpoint",
        ],
        kind="stable",
    ).reset_index(drop=True)
    change_columns = [
        "Coverage Mean Change (% of Total Population)",
        "Grid Probability P95 Change (percentage points)",
        "Hospital Catchment P95 Change (% of Total Population)",
    ]
    reference = table["Simulation Replicate Checkpoint"].eq(100)
    for column in change_columns:
        table[column] = table[column].astype(object).where(~reference, "None")
    table = table[TABLE_COLUMNS]
    if table.shape != (20, 10):
        raise RuntimeError(f"Unexpected convergence table shape: {table.shape}")
    if table.isna().any().any():
        raise RuntimeError("Convergence table contains missing values")
    return table


def style_workbook(path: Path, data_rows: int) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook["Convergence Summary"]
    navy = "1F4E78"
    main_fill = "F4F8FB"
    main_alt_fill = "E8F1F7"
    stress_fill = "FCE4D6"
    note_fill = "EAF3F8"
    stable_fill = "E2F0D9"
    unstable_fill = "F4CCCC"
    reference_fill = "E7E6E6"
    white = "FFFFFF"
    thin_grey = Side(style="thin", color="B7C3CA")
    medium_grey = Side(style="medium", color="7F8C8D")

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    title = sheet.cell(1, 1)
    title.value = "Monte Carlo Convergence Summary"
    title.font = Font(name="Arial", size=14, bold=True, color=white)
    title.fill = PatternFill("solid", fgColor=navy)
    title.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 24

    header_row = 3
    for cell in sheet[header_row]:
        cell.font = Font(name="Arial", size=8.5, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=thin_grey,
            bottom=thin_grey,
            left=thin_grey,
            right=thin_grey,
        )
    sheet.row_dimensions[header_row].height = 72

    first_data_row = 4
    last_data_row = first_data_row + data_rows - 1
    severity_group = -1
    previous_severity: float | None = None
    for row in range(first_data_row, last_data_row + 1):
        severity = float(sheet.cell(row, 1).value)
        starts_group = previous_severity is None or not np.isclose(
            severity, previous_severity
        )
        if starts_group:
            severity_group += 1
        if np.isclose(severity, 10.0):
            fill = PatternFill("solid", fgColor=stress_fill)
        else:
            fill = PatternFill(
                "solid", fgColor=main_fill if severity_group % 2 == 0 else main_alt_fill
            )
        top_side = medium_grey if starts_group and row > first_data_row else thin_grey
        for column in range(1, 11):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Arial", size=8.5)
            cell.fill = fill
            cell.border = Border(
                top=top_side,
                bottom=thin_grey,
                left=thin_grey,
                right=thin_grey,
            )
            cell.alignment = Alignment(
                horizontal="left" if column == 10 else "right",
                vertical="center",
            )
        status = str(sheet.cell(row, 10).value)
        status_fill = {
            "Stable": stable_fill,
            "Not stable": unstable_fill,
            "Reference checkpoint": reference_fill,
        }.get(status)
        if status_fill is None:
            raise RuntimeError(f"Unexpected convergence status in workbook: {status}")
        sheet.cell(row, 10).fill = PatternFill("solid", fgColor=status_fill)
        sheet.cell(row, 10).font = Font(name="Arial", size=8.5, bold=True)

        sheet.cell(row, 1).number_format = '0"%"'
        sheet.cell(row, 2).number_format = "#,##0"
        for column in (3, 5, 6):
            sheet.cell(row, column).number_format = "#,##0"
        sheet.cell(row, 4).number_format = "#,##0.0"
        for column in (7, 9):
            sheet.cell(row, column).number_format = '0.000"%"'
        sheet.cell(row, 8).number_format = '0.000" pp"'
        previous_severity = severity

    note_row = last_data_row + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
    note = sheet.cell(note_row, 1)
    note.value = (
        "Note: Coverage statistics refer to population retaining 30-minute emergency "
        "access. Each checkpoint uses the first R paired replicates. Change statistics "
        "compare with the preceding checkpoint; None at 100 replicates indicates that no "
        "previous checkpoint exists. Stable requires coverage mean change no greater than "
        "0.5% of total population, grid-probability P95 change no greater than 5 percentage "
        "points, and hospital-catchment P95 change no greater than 0.5% of total population. "
        "The 1%, 3%, and 5% states are main scenarios; 10% is an extreme stress scenario."
    )
    note.font = Font(name="Arial", size=9, italic=True, color="37474F")
    note.fill = PatternFill("solid", fgColor=note_fill)
    note.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[note_row].height = 62

    widths = [17, 18, 22, 20, 18, 18, 21, 22, 23, 23]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:J{last_data_row}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(path)


def main() -> None:
    frame = pd.read_parquet(INPUT)
    summary = json.loads((EXPERIMENT / "experiment_summary.json").read_text())
    total_population = float(summary["network"]["total_population"])
    validate_source(frame, total_population)
    table = build_table(frame, total_population)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        table.to_excel(
            writer, sheet_name="Convergence Summary", startrow=2, index=False
        )
    style_workbook(OUTPUT, len(table))
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Rows: {len(table)}; columns: {len(table.columns)}")
    print(table.to_string(index=False))


if __name__ == "__main__":
    main()
