#!/usr/bin/env python3
"""Emergency Access Reliability by Severity and Threshold.

Plan: Compare population-weighted emergency-access reliability across the three
main failure severities, the stress scenario, and three timely-access thresholds.
Framework: AnaSOP Sections 5.1-5.2, 6.2-6.3, and Analytical Workflow step 4.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
INPUT = (
    ROOT
    / "data"
    / "exp"
    / "monte_carlo_length_weighted_full_1000"
    / "replicate_metrics.parquet"
)
OUTPUT = (
    ROOT
    / "data"
    / "results"
    / "tables"
    / "Table_emergency_access_reliability_by_severity_and_threshold.xlsx"
)
THRESHOLDS = (15, 30, 45)
SEVERITIES = (0.01, 0.03, 0.05, 0.10)
COLUMNS = [
    "Expected Failed Road Length (%)",
    "Timely Access Threshold (min)",
    "Simulation Replicates",
    "Baseline Population with Timely Access",
    "Mean Population with Timely Access",
    "Coverage Monte Carlo SE",
    "Coverage P05",
    "Coverage P50",
    "Coverage P95",
    "Mean Population Losing Timely Access",
    "Mean Older Population Losing Timely Access",
    "Mean Population Newly Disconnected",
]


def summarize(metrics: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, float | int]] = []
    for severity in SEVERITIES:
        subset = metrics.loc[
            np.isclose(metrics["Expected Failed Road Length Share"], severity)
        ].copy()
        if len(subset) != 1000:
            raise RuntimeError(
                f"Expected 1,000 replicates at severity {severity:.2%}; got {len(subset)}"
            )
        if subset["Simulation Replicate"].nunique() != len(subset):
            raise RuntimeError(f"Simulation Replicate is not unique at {severity:.2%}")

        for threshold in THRESHOLDS:
            coverage_column = f"Population within {threshold} Minutes"
            loss_column = (
                f"Population Losing Baseline {threshold}-Minute Access"
            )
            older_loss_column = (
                f"Older Population Losing Baseline {threshold}-Minute Access"
            )
            coverage = subset[coverage_column].astype(float)
            loss = subset[loss_column].astype(float)
            older_loss = subset[older_loss_column].astype(float)
            newly_disconnected = subset["Population Newly Disconnected"].astype(float)

            baseline = coverage + loss
            if baseline.nunique() != 1:
                raise RuntimeError(
                    f"Baseline timely population varies at {severity:.2%}, "
                    f"threshold {threshold}"
                )
            if (loss < 0).any() or (older_loss < 0).any():
                raise RuntimeError("A timely-access loss measure is negative")

            replicates = len(coverage)
            rows.append(
                {
                    "Expected Failed Road Length (%)": 100.0 * severity,
                    "Timely Access Threshold (min)": threshold,
                    "Simulation Replicates": replicates,
                    "Baseline Population with Timely Access": float(baseline.iloc[0]),
                    "Mean Population with Timely Access": float(coverage.mean()),
                    "Coverage Monte Carlo SE": float(
                        coverage.std(ddof=1) / np.sqrt(replicates)
                    ),
                    "Coverage P05": float(coverage.quantile(0.05)),
                    "Coverage P50": float(coverage.quantile(0.50)),
                    "Coverage P95": float(coverage.quantile(0.95)),
                    "Mean Population Losing Timely Access": float(loss.mean()),
                    "Mean Older Population Losing Timely Access": float(
                        older_loss.mean()
                    ),
                    "Mean Population Newly Disconnected": float(
                        newly_disconnected.mean()
                    ),
                }
            )

    table = pd.DataFrame(rows, columns=COLUMNS)
    if table.shape != (12, 12):
        raise RuntimeError(f"Unexpected table shape: {table.shape}")
    if table.isna().any().any():
        raise RuntimeError("The reliability table contains missing values")
    return table


def style_workbook(path: Path, data_rows: int) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook["Reliability Summary"]
    navy = "1F4E78"
    main_fill = "DCE6F1"
    stress_fill = "FCE4D6"
    note_fill = "EAF3F8"
    white = "FFFFFF"
    thin_grey = Side(style="thin", color="B7C3CA")
    medium_grey = Side(style="medium", color="7F8C8D")

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
    title = sheet.cell(1, 1)
    title.value = "Emergency Access Reliability by Severity and Threshold"
    title.font = Font(name="Arial", size=14, bold=True, color=white)
    title.fill = PatternFill("solid", fgColor=navy)
    title.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 24

    header_row = 3
    for cell in sheet[header_row]:
        cell.font = Font(name="Arial", size=8.5, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=thin_grey,
            bottom=thin_grey,
            left=thin_grey,
            right=thin_grey,
        )
    sheet.row_dimensions[header_row].height = 58

    first_data_row = 4
    last_data_row = first_data_row + data_rows - 1
    for row in range(first_data_row, last_data_row + 1):
        severity = float(sheet.cell(row, 1).value)
        fill = PatternFill(
            "solid", fgColor=stress_fill if np.isclose(severity, 10.0) else main_fill
        )
        starts_group = row == first_data_row or not np.isclose(
            severity, float(sheet.cell(row - 1, 1).value)
        )
        top_side = medium_grey if starts_group and row > first_data_row else thin_grey
        for column in range(1, 13):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Arial", size=8.5)
            cell.fill = fill
            cell.border = Border(
                top=top_side,
                bottom=thin_grey,
                left=thin_grey,
                right=thin_grey,
            )
            cell.alignment = Alignment(horizontal="right", vertical="center")

        sheet.cell(row, 1).number_format = '0"%"'
        sheet.cell(row, 2).number_format = '0" min"'
        sheet.cell(row, 3).number_format = "#,##0"
        for column in range(4, 13):
            sheet.cell(row, column).number_format = "#,##0"
        sheet.cell(row, 6).number_format = "#,##0.0"

    note_row = last_data_row + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=12)
    note = sheet.cell(note_row, 1)
    note.value = (
        "Note: The 1%, 3%, and 5% states are the main scenarios; 10% is an extreme "
        "stress scenario. Statistics are calculated across 1,000 paired Monte Carlo "
        "replicates. P05, P50, and P95 describe the distribution of population retaining "
        "timely access. Older-population loss uses its valid disclosure-group support. "
        "New disconnection is threshold-independent and is repeated across thresholds "
        "within each severity for a complete severity-threshold record."
    )
    note.font = Font(name="Arial", size=9, italic=True, color="37474F")
    note.fill = PatternFill("solid", fgColor=note_fill)
    note.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[note_row].height = 50

    widths = [
        17,
        17,
        15,
        21,
        21,
        17,
        15,
        15,
        15,
        22,
        24,
        22,
    ]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:L{last_data_row}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(path)


def main() -> None:
    metrics = pd.read_parquet(INPUT)
    required = {
        "Simulation Replicate",
        "Expected Failed Road Length Share",
        "Population Newly Disconnected",
    }
    for threshold in THRESHOLDS:
        required.update(
            {
                f"Population within {threshold} Minutes",
                f"Population Losing Baseline {threshold}-Minute Access",
                f"Older Population Losing Baseline {threshold}-Minute Access",
            }
        )
    missing = required.difference(metrics.columns)
    if missing:
        raise RuntimeError(f"Missing required columns: {sorted(missing)}")

    table = summarize(metrics)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Reliability Summary", startrow=2, index=False)
    style_workbook(OUTPUT, len(table))
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Rows: {len(table)}; columns: {len(table.columns)}")
    print(table.to_string(index=False))


if __name__ == "__main__":
    main()
